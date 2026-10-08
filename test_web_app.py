import os
import tempfile
import time
import unittest
from unittest.mock import patch

from kaliper_app import KaliperStandalone
import qc_queries
from web_app import create_app


class FakeDB:
    def __init__(self):
        self.rows, self.fail = [], False

    def insert_line(self, line, tipe, mold, sides, avg, ts):
        if self.fail:
            raise RuntimeError("DB mati (simulasi)")
        self.rows.append((line, tipe, mold, sides, avg))

    def get_tipe(self, mold):
        return {"M-01": "GRID A"}.get(mold.upper())

    def connect(self):
        raise RuntimeError("tidak dipakai di tes")


def wait_for(cond, timeout=2.0):
    t0 = time.time()
    while time.time() - t0 < timeout:
        if cond():
            return True
        time.sleep(0.02)
    return False


class WebTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.db = FakeDB()
        core = KaliperStandalone(filename=os.path.join(self.tmp, "t.csv"), db=self.db)
        self.app = create_app(core)
        self.c = self.app.test_client()

    def ctx(self, mold="m-01", tipe="grid a"):
        return self.c.post("/api/context", json={"mold": mold, "tipe": tipe})

    def isi(self, vals):
        for v in vals:
            r = self.c.post("/api/input", json={"value": v})
            self.assertEqual(r.status_code, 200, r.json)

    def test_halaman_utama_dan_aset(self):
        response = self.c.get("/")
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn('id="lineSelectMain"', page)
        self.assertIn('id="val-tipe-main"', page)
        self.assertIn('id="val-mold-main"', page)
        self.assertIn('id="pointGrid"', page)
        self.assertIn('id="caliperInput" class="caliper-capture-input"', page)
        self.assertNotIn("Input Kaliper (tembak kaliper", page)
        for f in ("chart.min.js", "hammer.min.js", "chartjs-plugin-zoom.min.js", "Logo GS.png"):
            self.assertEqual(self.c.get("/static/" + f).status_code, 200, f)

    def test_input_ditolak_sebelum_konfirmasi_mold(self):
        r = self.c.post("/api/input", json={"value": "1.0"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.c.get("/api/state").json["count"], 0)

    def test_8_sisi_masuk_db_dan_pindah_line(self):
        self.assertEqual(self.ctx().status_code, 200)
        self.isi(["1,10", "1,20", "1,30", "1,40", "1,50", "1,60", "1,70", "1,80"])
        self.assertTrue(wait_for(lambda: len(self.db.rows) == 1))
        line, tipe, mold, sides, avg = self.db.rows[0]
        self.assertEqual((line, tipe, mold), ("C01", "GRID A", "M-01"))
        self.assertAlmostEqual(avg, 1.45)
        st = self.c.get("/api/state").json
        self.assertEqual(st["active_line"], "C02")
        self.assertEqual(st["mold"], "")                 # konteks dikosongkan -> operator konfirmasi lagi
        self.assertEqual(st["last_mold"], "M-01")
        self.assertEqual(st["last_complete"]["line"], "C01")
        self.assertTrue(wait_for(lambda: self.c.get("/api/state").json["lines"]["C01"]["state"] == "ok"))

    def test_input_tidak_valid_dan_undo(self):
        self.ctx()
        self.assertEqual(self.c.post("/api/input", json={"value": "abc"}).status_code, 400)
        self.isi(["1.0", "2.0"])
        r = self.c.post("/api/undo")
        self.assertEqual(r.json["state"]["count"], 1)
        self.assertEqual(r.json["state"]["titik1"], 1.0)
        self.assertEqual(r.json["state"]["titik2"], "-")

    def test_db_gagal_ditahan_lalu_kirim_ulang(self):
        self.db.fail = True
        self.ctx()
        self.isi(["1.0"] * 8)
        self.assertTrue(wait_for(lambda: self.c.get("/api/state").json["lines"]["C01"]["state"] == "fail"))
        self.assertEqual(self.c.get("/api/state").json["pending"], 1)
        self.db.fail = False
        self.c.post("/api/flush")
        self.assertTrue(wait_for(lambda: len(self.db.rows) == 1))
        self.assertTrue(wait_for(lambda: self.c.get("/api/state").json["pending"] == 0))

    def test_pindah_line_dan_ganti_mold_diblok_saat_ada_nilai(self):
        self.ctx()
        self.isi(["1.0"])
        self.assertEqual(self.c.post("/api/goto_line", json={"line": "C05"}).status_code, 400)
        self.assertEqual(self.ctx("m-02", "grid b").status_code, 400)
        self.c.post("/api/clear_line")
        self.assertEqual(self.c.post("/api/goto_line", json={"line": "C05"}).status_code, 200)
        self.assertEqual(self.c.get("/api/state").json["active_line"], "C05")
        self.assertEqual(self.c.post("/api/goto_line", json={"line": "C99"}).status_code, 400)

    def test_mode_csv_saja_tanpa_db(self):
        core = KaliperStandalone(filename=os.path.join(self.tmp, "x.csv"), db=None)
        c = create_app(core).test_client()
        c.post("/api/context", json={"mold": "m", "tipe": "t"})
        for _ in range(8):
            c.post("/api/input", json={"value": "1.5"})
        st = c.get("/api/state").json
        self.assertEqual(st["db_server"], "csv only")
        self.assertEqual(st["pending"], 1)
        self.assertEqual(c.get("/api/history?mold=m").json["error"], True)

    def test_query_history_memakai_sintaks_sql_server(self):
        sql, params = qc_queries._union(
            "C02", "M-01", "GRID A", "2025-01-01", None, None, None)
        self.assertIn("FROM dbo.tbl_c02", sql)
        self.assertIn("[timestamp] AS ts", sql)
        self.assertIn("LTRIM(RTRIM(no_mold))", sql)
        self.assertEqual(sql.count("?"), 3)
        self.assertEqual(params, ["M-01", "GRID A", "2025-01-01"])
        self.assertNotIn("%s", sql)

    def test_export_excel_header_sama_dengan_frontend(self):
        import io
        from openpyxl import load_workbook
        rows = [
            {"line": "C01", "mold": "M-01", "tipe": "GRID A", "date": "2026-10-08", "time": "13:41:24",
             **{f"titik{i}": 1.8 + i / 100 for i in range(1, 9)}, "avg": 1.845},
            {"line": "C02", "mold": "M-01", "tipe": "GRID A", "date": "2026-10-08", "time": "14:00:00",
             **{f"titik{i}": None for i in range(1, 9)}, "avg": None},
        ]
        with patch("qc_queries.history", return_value={"rows": rows, "tipe": "GRID A"}):
            self.app.kal.core.db = self.db
            r = self.c.get("/api/export_history?mold=M-01")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.mimetype)
        self.assertIn("QC_History_M-01_", r.headers["Content-Disposition"])
        ws = load_workbook(io.BytesIO(r.data))["QC History"]
        self.assertEqual([ws.cell(1, c).value for c in (1, 2, 3, 4, 5, 13)],
                         ["LINE", "NO. MOLD", "TANGGAL", "JAM", "TITIK", "AVERAGE"])
        self.assertEqual([ws.cell(2, c).value for c in range(5, 13)], list(range(1, 9)))
        self.assertEqual(ws.freeze_panes, "A3")
        self.assertEqual((ws["A3"].value, ws["B3"].value, ws["M3"].value), ("C01", "M-01", 1.845))
        self.assertIsNone(ws["E4"].value)
        self.assertEqual(ws["E5"].value, '=IF(COUNT(E3:E4)=0,"-",MIN(E3:E4))')

    def test_export_ditolak_tanpa_filter_atau_tanpa_db(self):
        with patch("qc_queries.history", return_value={"rows": [], "tipe": "-"}):
            self.assertEqual(self.c.get("/api/export_history").status_code, 404)
        core = KaliperStandalone(filename=os.path.join(self.tmp, "y.csv"), db=None)
        self.assertEqual(create_app(core).test_client().get("/api/export_history?mold=x").status_code, 400)


if __name__ == "__main__":
    unittest.main()
