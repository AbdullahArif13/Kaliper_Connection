"""
Ekspor riwayat pengukuran ke Excel (.xlsx) dengan header yang sama seperti tabel di frontend:

    Line | No. Mold | Tanggal | Jam | Titik (1..8) | AVERAGE      <- baris 1-2 (header)
    ... data (terbaru dulu, sama seperti tampilan) ...
    Min / Max / AVERAGE per kolom                               <- rumus Excel di bawah data

Sheet kedua "Filter" mencatat filter & waktu ekspor supaya file bisa ditelusuri.
"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SISI = 8
FONT = "Arial"
NUM_FMT = "0.000"
HEADER_ROWS = 2
FIRST_DATA_ROW = HEADER_ROWS + 1
LAST_COL = 4 + SISI + 1            # A..D + 8 titik + AVERAGE = 13 (kolom M)

_thin = Side(style="thin", color="C9D1DE")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CENTER = Alignment(horizontal="center", vertical="center")
GROUP_FILL = PatternFill("solid", fgColor="EEF4FF")     # .group-head di frontend
SUB_FILL = PatternFill("solid", fgColor="F8FBFF")       # th biasa di frontend
FOOT_FILL = PatternFill("solid", fgColor="FBFDFF")      # .footer-summary-row di frontend
GROUP_FONT = Font(name=FONT, bold=True, size=10, color="003A8F")
SUB_FONT = Font(name=FONT, bold=True, size=10, color="0D1D36")
BODY_FONT = Font(name=FONT, size=10)
MOLD_FONT = Font(name=FONT, size=10, bold=True, color="0D5FE9")


def _num(v):
    try:
        return float(v) if v is not None and v != "" and v != "-" else None
    except (TypeError, ValueError):
        return None


def _date(s):
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return s if s not in (None, "-") else None


def _time(s):
    try:
        return datetime.strptime(s, "%H:%M:%S").time()
    except (TypeError, ValueError):
        return s if s not in (None, "-") else None


def _cell(ws, row, col, value=None, font=BODY_FONT, fill=None, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.alignment, c.border = font, CENTER, BORDER
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def build_history_workbook(rows, filters, truncated=False):
    """rows = daftar dict dari qc_queries.history()['rows']; filters = dict filter yang dipakai."""
    wb = Workbook()
    ws = wb.active
    ws.title = "QC History"

    # ---- header 2 baris (rowspan/colspan seperti <thead> di frontend) ----
    for col, label in enumerate(["LINE", "NO. MOLD", "TANGGAL", "JAM"], start=1):
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
        for r in (1, 2):                                   # border pada seluruh sel gabungan
            _cell(ws, r, col, label if r == 1 else None, GROUP_FONT, GROUP_FILL)
    ws.merge_cells(start_row=1, start_column=5, end_row=1, end_column=4 + SISI)
    for col in range(5, 5 + SISI):
        _cell(ws, 1, col, "TITIK" if col == 5 else None, GROUP_FONT, GROUP_FILL)
        _cell(ws, 2, col, col - 4, SUB_FONT, SUB_FILL)
    ws.merge_cells(start_row=1, start_column=LAST_COL, end_row=2, end_column=LAST_COL)
    for r in (1, 2):
        _cell(ws, r, LAST_COL, "AVERAGE" if r == 1 else None, GROUP_FONT, GROUP_FILL)
    # ---- data ----
    r = FIRST_DATA_ROW
    for row in rows:
        _cell(ws, r, 1, row.get("line"))
        _cell(ws, r, 2, row.get("mold"), MOLD_FONT)
        _cell(ws, r, 3, _date(row.get("date")), fmt="yyyy-mm-dd")
        _cell(ws, r, 4, _time(row.get("time")), fmt="hh:mm:ss")
        for i in range(SISI):
            _cell(ws, r, 5 + i, _num(row.get(f"titik{i + 1}")), fmt=NUM_FMT)
        _cell(ws, r, LAST_COL, _num(row.get("avg")), fmt=NUM_FMT)
        r += 1
    last = r - 1

    # ---- ringkasan Min / Max / AVERAGE (rumus, sama seperti <tfoot> di frontend) ----
    for label, fn in (("MIN", "MIN"), ("MAX", "MAX"), ("AVERAGE", "AVERAGE")):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        for col in range(1, 5):
            _cell(ws, r, col, label if col == 1 else None, GROUP_FONT, FOOT_FILL)
        for col in range(5, LAST_COL + 1):
            rng = f"{get_column_letter(col)}{FIRST_DATA_ROW}:{get_column_letter(col)}{last}"
            _cell(ws, r, col, f'=IF(COUNT({rng})=0,"-",{fn}({rng}))', SUB_FONT, FOOT_FILL, NUM_FMT)
        r += 1

    # ---- tampilan ----
    for col, w in zip(range(1, LAST_COL + 1), [8, 16, 12, 10] + [9] * SISI + [12]):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 20
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)       # header tetap terlihat saat scroll
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"1:{HEADER_ROWS}"
    wb.calculation.fullCalcOnLoad = True                           # Excel menghitung ulang rumus saat dibuka

    # ---- sheet info filter ----
    info = wb.create_sheet("Filter")
    items = [
        ("Line", filters.get("line") or "Semua"), ("No. Mold", filters.get("mold") or "Semua"),
        ("Tipe", filters.get("tipe") or "Semua"),
        ("Tanggal", f"{filters.get('date_from') or '-'} s/d {filters.get('date_to') or '-'}"),
        ("Jam", f"{filters.get('time_from') or '-'} s/d {filters.get('time_to') or '-'}"),
        ("Jumlah baris", len(rows)), ("Diekspor pada", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    if truncated:
        items.append(("Catatan", "Data dibatasi jumlah maksimum ekspor; persempit filter untuk data lengkap."))
    for i, (k, v) in enumerate(items, start=1):
        a, b = info.cell(row=i, column=1, value=k), info.cell(row=i, column=2, value=v)
        a.font, b.font = Font(name=FONT, bold=True, size=10), Font(name=FONT, size=10)
        b.alignment = Alignment(horizontal="left")
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 16, 48

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
